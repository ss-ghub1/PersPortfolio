"""
Loads a UBS position-snapshot export (one tab per portfolio, blocks of rows
grouped by 'Group of products', with a repeating header row per block and a
footer metadata block) into position_snapshot.

The two header variants (cash-account rows vs security rows) share the same
leading 23 columns and the same trailing 4 columns (market value / % / accrued
interest / % accrued) - only 5 columns in between differ, which is what this
parser special-cases.
"""
import re
import sqlite3
from datetime import datetime
from pathlib import Path

import openpyxl

from db import get_connection, _clean_account_id, resolve_account_id, is_source_file_loaded

COMMON_LEN = 23  # columns 0..22 identical in both header variants


def _parse_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    s = str(value).strip()
    m = re.match(r"(\d{2})\.(\d{2})\.(\d{4})", s)
    if m:
        d, mth, y = m.groups()
        return f"{y}-{mth}-{d}"
    return s


def _to_float(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).replace("'", "").replace(",", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def _upsert_instrument(conn, isin, valor, name, asset_class, ccy, sector):
    instrument_id = isin or (f"VALOR:{valor}" if valor else None)
    if not instrument_id:
        return None
    conn.execute(
        """INSERT INTO instrument (instrument_id, isin, valor, name, asset_class, instrument_currency, sector)
           VALUES (?, ?, ?, ?, ?, ?, ?)
           ON CONFLICT(instrument_id) DO UPDATE SET
             name=excluded.name,
             asset_class=COALESCE(instrument.asset_class, excluded.asset_class),
             instrument_currency=COALESCE(instrument.instrument_currency, excluded.instrument_currency),
             sector=COALESCE(instrument.sector, excluded.sector)
        """,
        (instrument_id, isin, str(valor) if valor else None, name, asset_class, ccy, sector),
    )
    return instrument_id


def parse_position_file(path: Path):
    """Yields one dict per position row across all sheets."""
    wb = openpyxl.load_workbook(path, data_only=True)
    source_file = Path(path).name

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=32, values_only=True))

        # First pass: footer metadata lives at the bottom of each sheet but
        # applies to every row above it, so read it before processing rows.
        as_of_date = None
        valued_in = None
        native_account_code = None
        for row in rows:
            cell0 = row[0]
            if isinstance(cell0, str) and cell0.startswith("Portfolio number:"):
                native_account_code = cell0.split(":", 1)[1].strip()
            elif isinstance(cell0, str) and cell0.startswith("Valued as of:"):
                as_of_date = _parse_date(cell0.split(":", 1)[1].strip())
            elif isinstance(cell0, str) and cell0.startswith("Valued in:"):
                valued_in = cell0.split(":", 1)[1].strip()

        block_is_cash = None
        i = 0
        while i < len(rows):
            row = rows[i]
            cell0 = row[0]

            if cell0 == "Banking relationship":
                # header row -> determine block type from column index 23
                block_is_cash = (row[23] == "Account no.")
                i += 1
                continue

            if cell0 == "Exchange rates" or cell0 is None:
                i += 1
                continue

            if isinstance(cell0, str) and (cell0.startswith("Portfolio number:")
                                            or cell0.startswith("Valued as of:")
                                            or cell0.startswith("Valued in:")
                                            or cell0.startswith("Export created on:")
                                            or cell0.startswith("Number of positions:")):
                i += 1
                continue

            # FX rate rows inside "Exchange rates" sub-block: e.g.
            # ('EUR/SGD', 1.46513, 'SGD/EUR', 0.682533, ...) - capture, don't just skip
            if isinstance(cell0, str) and "/" in cell0 and len(cell0) <= 8 and row[1] is not None:
                for pair_idx in (0, 2):
                    pair, rate = row[pair_idx], row[pair_idx + 1]
                    if isinstance(pair, str) and "/" in pair and isinstance(rate, (int, float)):
                        from_ccy, to_ccy = pair.split("/")
                        _FX_RATES.append((as_of_date, from_ccy, to_ccy, float(rate), source_file))
                i += 1
                continue

            if block_is_cash is None or row[1] is None:
                i += 1
                continue

            # --- a genuine data row ---
            group_of_products = row[2]
            is_cash = (group_of_products == "Liquidity - Accounts")
            ccy = row[4]
            isin = row[8]
            valor = row[7]
            name = row[13]
            sector = row[11]

            if block_is_cash:
                market_value_base = _to_float(row[28])
                pct_mv = _to_float(row[29])
                accrued_interest = _to_float(row[30])
                cost_value = None
                market_price = None
            else:
                cost_value = _to_float(row[24])
                market_price = _to_float(row[25])
                market_value_base = _to_float(row[28])
                pct_mv = _to_float(row[29])
                accrued_interest = _to_float(row[30])

            yield {
                "native_account_code": row[1],
                "instrument_id": None if is_cash else _upsert_instrument(
                    _CONN, isin, valor, name, group_of_products, ccy, sector
                ),
                "is_cash": is_cash,
                "as_of_date": as_of_date,
                "position_currency": ccy,
                "quantity": _to_float(row[5]),
                "cost_price": _to_float(row[9]),
                "cost_value_native": cost_value,
                "market_price": market_price,
                "market_value_native": None,  # UBS export only gives market value in base ccy directly
                "market_value_base": market_value_base,
                "unrealized_pl_pct": _to_float(row[27]) if not block_is_cash else None,
                "accrued_interest": accrued_interest,
                "pct_of_portfolio": pct_mv,
                "source_file": source_file,
            }
            i += 1


_CONN = None  # set by load_positions so the generator can upsert instruments inline
_FX_RATES = []  # (as_of_date, from_ccy, to_ccy, rate, source_file) collected during parsing


def load_positions(path: Path, conn: sqlite3.Connection = None, force: bool = False):
    global _CONN, _FX_RATES
    own_conn = conn is None
    conn = conn or get_connection()

    source_file = Path(path).name
    if not force and is_source_file_loaded(conn, "position_snapshot", source_file):
        print(f"[positions] SKIPPED: '{source_file}' already loaded (use --force to reload anyway)")
        if own_conn:
            conn.close()
        return 0

    _CONN = conn
    _FX_RATES = []

    # Resolve account_id for every record first (single pass), so we know
    # exactly which (account_id, as_of_date) snapshots this load will touch
    # BEFORE inserting anything.
    n_skipped_no_account = 0
    records = []
    for rec in parse_position_file(path):
        account_id = resolve_account_id(conn, "UBS", rec["native_account_code"], alias_type="PORTFOLIO_CODE")
        if not account_id:
            n_skipped_no_account += 1
            continue
        rec["account_id"] = account_id
        records.append(rec)

    # Idempotency: clear any existing snapshot for each (account_id, as_of_date)
    # this load covers before inserting, so accidentally running the same load
    # twice replaces the snapshot instead of doubling every position - this
    # bit us for real (see commit history) when a snapshot got loaded twice
    # and every account total came out exactly 2x.
    touched = {(r["account_id"], r["as_of_date"]) for r in records}
    for account_id, as_of_date in touched:
        conn.execute(
            "DELETE FROM position_snapshot WHERE account_id=? AND as_of_date=?",
            (account_id, as_of_date),
        )

    n_loaded = 0
    for rec in records:
        conn.execute(
            """INSERT INTO position_snapshot
                 (account_id, instrument_id, is_cash, as_of_date, position_currency,
                  quantity, cost_price, cost_value_native, market_price,
                  market_value_native, market_value_base, unrealized_pl_pct,
                  accrued_interest, pct_of_portfolio, source_file)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                rec["account_id"], rec["instrument_id"], int(rec["is_cash"]), rec["as_of_date"],
                rec["position_currency"], rec["quantity"], rec["cost_price"],
                rec["cost_value_native"], rec["market_price"], rec["market_value_native"],
                rec["market_value_base"], rec["unrealized_pl_pct"], rec["accrued_interest"],
                rec["pct_of_portfolio"], rec["source_file"],
            ),
        )
        n_loaded += 1
    conn.commit()
    print(f"[positions] cleared {len(touched)} existing (account, as_of_date) snapshot(s) before reload")
    print(f"[positions] loaded {n_loaded} rows, skipped {n_skipped_no_account} "
          f"(unknown account_id - add to config/accounts.csv)")

    n_fx = 0
    for as_of_date, from_ccy, to_ccy, rate, source_file in _FX_RATES:
        if not as_of_date:
            continue
        conn.execute(
            """INSERT INTO fx_rate (rate_date, from_ccy, to_ccy, rate, source_file)
               VALUES (?, ?, ?, ?, ?)
               ON CONFLICT(rate_date, from_ccy, to_ccy) DO UPDATE SET rate=excluded.rate""",
            (as_of_date, from_ccy, to_ccy, rate, source_file),
        )
        n_fx += 1
    conn.commit()
    print(f"[positions] captured {n_fx} FX rate rows")

    if own_conn:
        conn.close()
    return n_loaded


if __name__ == "__main__":
    import sys
    force = "--force" in sys.argv
    load_positions(Path(sys.argv[1]), force=force)
