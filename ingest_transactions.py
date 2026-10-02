"""
Loads a UBS transaction extract workbook into txn. The workbook has two
different tab formats:

  - 'UBSPortTxns': one row per security-side event (trade/income) across all
    portfolios, identified by 'Portfolio' column directly.
  - per-currency cash ledger tabs (e.g. 'UBSAcctTxnsUSD', 'SGD1', 'EUR1'...):
    one row per cash movement in a single currency sub-account, with a
    running balance. These tabs identify their account only by a raw cash
    sub-account number in the header metadata - resolved via account_alias.

Both legs are classified through the shared rule table (classify.py) and
loaded into the same txn table, tagged by source_leg.
"""
import re
import sqlite3
from datetime import datetime, date
from pathlib import Path

import openpyxl

from db import get_connection, resolve_account_id, is_source_file_loaded
from classify import classify
from ingest_positions import _upsert_instrument  # reuse instrument upsert


def _normalize_ref(value):
    """Reference numbers can arrive as a zero-padded string (CSV) or as a
    number with the leading zero stripped (Excel reads it as numeric) -
    normalize both to the same canonical string so the dedup key matches
    regardless of which format a given extract used."""
    if value is None:
        return None
    s = str(value).strip()
    if s.lstrip("-").isdigit():
        return str(int(s))
    return s


def _to_float(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    s = s.replace("'", "").replace(",", "")
    s = re.sub(r"\s*p$", "", s)  # UBS sometimes suffixes share quantities with ' p'
    try:
        return float(s)
    except ValueError:
        return None


def _to_iso_date(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    s = str(value).strip()
    m = re.match(r"(\d{2})\.(\d{2})\.(\d{4})", s)
    if m:
        d, mth, y = m.groups()
        return f"{y}-{mth}-{d}"
    return s


# ---------------------------------------------------------------------------
# Security-leg tab: 'UBSPortTxns' (same column layout arrives as either an
# .xlsx tab or a standalone .csv export - 'Transaction list: All
# transactions...' style export). Both are parsed through this shared
# per-row extractor so a CSV pull doesn't need its own field-mapping logic.
# ---------------------------------------------------------------------------
def _extract_security_leg_record(row, source_file, conn):
    if row[1] is None:  # 'Banking relationship' blank -> not a data row
        return None
    native_account_code = row[2]
    trade_date = _to_iso_date(row[4])
    booking_date = _to_iso_date(row[6])
    value_date = _to_iso_date(row[7])
    raw_label = row[8]  # 'Description 1' e.g. 'Dividend', 'Stock Market Spot Purchase'
    counterparty_or_security = row[9]
    isin = row[12]
    valor = row[11]
    native_ccy = row[15]
    quantity = _to_float(row[14])
    trans_price = _to_float(row[16])
    fx_rate = _to_float(row[17])
    valuation_ccy = row[18]
    trans_value = _to_float(row[19])
    realized_pl = _to_float(row[22])
    order_no = row[23]
    external_ref = row[24]
    asset_class = row[25]

    txn_type, txn_subtype = classify(conn, raw_label)

    instrument_id = _upsert_instrument(
        conn, isin, valor, counterparty_or_security, asset_class, native_ccy, None
    ) if (isin or valor) else None

    return {
        "native_account_code": native_account_code,
        "instrument_id": instrument_id,
        "source_leg": "SECURITY",
        "trade_date": trade_date,
        "booking_date": booking_date,
        "value_date": value_date,
        "txn_type": txn_type,
        "txn_subtype": txn_subtype,
        "currency": valuation_ccy or native_ccy,
        "quantity": quantity,
        "price": trans_price,
        "gross_amount": trans_value,
        "fee_amount": None,
        "tax_amount": None,
        "running_balance": None,
        "fx_rate": fx_rate,
        "realized_pl": realized_pl,
        "raw_txn_type_label": raw_label,
        "raw_description": counterparty_or_security,
        "source_reference": _normalize_ref(external_ref or order_no),
        "source_file": source_file,
    }


def _parse_security_leg_sheet(ws, source_file, conn):
    rows = list(ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=29, values_only=True))
    for row in rows:
        rec = _extract_security_leg_record(row, source_file, conn)
        if rec:
            yield rec


def _parse_security_leg_csv(path: Path, conn):
    import csv
    source_file = Path(path).name
    with open(path, encoding="utf-8-sig") as f:
        reader = csv.reader(f, delimiter=";")
        header = next(reader)
        for row in reader:
            if not row or not row[0] or row[0].startswith("Transaction list"):
                continue
            rec = _extract_security_leg_record(row, source_file, conn)
            if rec:
                yield rec


# ---------------------------------------------------------------------------
# Cash-leg tabs: per-currency ledgers with header metadata + running balance
# ---------------------------------------------------------------------------
def _parse_cash_leg_sheet(ws, source_file, conn):
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=14, values_only=True))

    raw_account_number, opening_balance, closing_balance, valued_in = None, None, None, None
    header_row_idx = None
    for idx, row in enumerate(rows):
        if row[0] == "Account number:":
            raw_account_number = str(row[1]).strip()
        elif row[0] == "Opening balance:":
            opening_balance = _to_float(row[1])
        elif row[0] == "Closing balance:":
            closing_balance = _to_float(row[1])
        elif row[0] == "Valued in:":
            valued_in = row[1]
        elif row[0] == "Trade date":
            header_row_idx = idx
            break

    if header_row_idx is None or raw_account_number is None:
        return

    meta = {
        "raw_account_number": raw_account_number,
        "opening_balance": opening_balance,
        "closing_balance": closing_balance,
        "currency": valued_in,
    }

    for row in rows[header_row_idx + 1:]:
        trade_date = row[0]
        if trade_date is None or row[10] == "Balance brought forward":
            continue
        booking_date = row[2]
        value_date = row[3]
        ccy = row[4]
        debit = _to_float(row[5])
        credit = _to_float(row[6])
        running_balance = _to_float(row[8])
        txn_no = row[9]
        description1 = row[10]  # usually counterparty/security name, but sometimes IS the type
        description2 = row[11]  # usually the type label, but blank for fee-style entries
        detail = row[12]

        # UBS's 'Debit' column already stores a negative number - do not negate again.
        gross_amount = credit if credit is not None else debit
        # UBS is inconsistent about which of the two description fields carries
        # the classifiable label, so match against both, preferring description2.
        classify_text = " | ".join(str(x) for x in (description2, description1) if x)
        txn_type, txn_subtype = classify(conn, classify_text)
        raw_label = description2 or description1
        counterparty_or_security = description1 if description2 else None

        # try to pull quantity/price back out of the free-text detail field,
        # e.g. "Number/Amt. 3896.095, Transaction price: 9.932499 USD"
        qty, price = None, None
        if detail:
            m_qty = re.search(r"Number/Amt\.\s*([\d'.,]+)", str(detail))
            m_price = re.search(r"Transaction price:\s*([\d'.,]+)", str(detail))
            if m_qty:
                qty = _to_float(m_qty.group(1))
            if m_price:
                price = _to_float(m_price.group(1))

        yield {
            "meta": meta,
            "instrument_id": None,  # cash leg does not reliably carry ISIN; security leg is the source of truth for holdings
            "source_leg": "CASH",
            "trade_date": _to_iso_date(trade_date),
            "booking_date": _to_iso_date(booking_date),
            "value_date": _to_iso_date(value_date),
            "txn_type": txn_type,
            "txn_subtype": txn_subtype,
            "currency": ccy,
            "quantity": qty,
            "price": price,
            "gross_amount": gross_amount,
            "fee_amount": None,
            "tax_amount": None,
            "running_balance": running_balance,
            "fx_rate": None,
            "realized_pl": None,
            "raw_txn_type_label": raw_label,
            "raw_description": counterparty_or_security,
            "source_reference": _normalize_ref(txn_no),
            "source_file": source_file,
        }


# ---------------------------------------------------------------------------
def load_transactions(path: Path, conn: sqlite3.Connection = None, force: bool = False):
    own_conn = conn is None
    conn = conn or get_connection()
    path = Path(path)

    source_file = path.name
    if not force and is_source_file_loaded(conn, "txn", source_file):
        print(f"[transactions] SKIPPED: '{source_file}' already loaded (use --force to reload anyway)")
        if own_conn:
            conn.close()
        return 0

    n_loaded, n_dupe_or_error, n_skipped_no_account = 0, 0, 0

    if path.suffix.lower() == ".csv":
        # standalone security-leg CSV export ('Transaction list: All
        # transactions...' style, same column layout as the UBSPortTxns tab)
        for rec in _parse_security_leg_csv(path, conn):
            account_id = resolve_account_id(conn, "UBS", rec["native_account_code"])
            if not account_id:
                n_skipped_no_account += 1
                continue
            n_loaded += _insert_txn(conn, account_id, rec)
        conn.commit()
        print(f"[transactions] loaded {n_loaded} rows from CSV, skipped {n_skipped_no_account} "
              f"(unresolved account)")
        if own_conn:
            conn.close()
        return n_loaded

    wb = openpyxl.load_workbook(path, data_only=True)
    source_file = Path(path).name
    n_unmatched_sheets = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_a1 = ws["A1"].value

        if header_a1 == "Valuation date":
            records = _parse_security_leg_sheet(ws, source_file, conn)
            for rec in records:
                account_id = resolve_account_id(conn, "UBS", rec["native_account_code"])
                if not account_id:
                    n_skipped_no_account += 1
                    continue
                n_loaded += _insert_txn(conn, account_id, rec)

        elif header_a1 == "Account number:":
            records = list(_parse_cash_leg_sheet(ws, source_file, conn))
            if not records:
                continue
            meta = records[0]["meta"]
            account_id = resolve_account_id(conn, "UBS", meta["raw_account_number"],
                                             alias_type="CASH_SUBACCOUNT")
            if not account_id:
                print(f"[transactions] WARNING: no account_alias mapping for cash "
                      f"sub-account '{meta['raw_account_number']}' (sheet '{sheet_name}') "
                      f"- add it to config/account_alias.csv. Skipped {len(records)} rows.")
                n_skipped_no_account += len(records)
                continue
            for rec in records:
                n_loaded += _insert_txn(conn, account_id, rec)

        else:
            n_unmatched_sheets += 1
            print(f"[transactions] WARNING: sheet '{sheet_name}' did not match a known "
                  f"tab format (A1='{header_a1}') - skipped.")

    if n_unmatched_sheets == len(wb.sheetnames):
        # EVERY sheet in the file was unrecognized - this is the signature
        # of the wrong file entirely (e.g. the position snapshot passed as
        # --transactions by mistake), not a benign extra tab. That incident
        # is exactly why this check exists: it completed "successfully"
        # with 0 rows loaded and only warnings printed, easy to miss in a
        # long command's output. A partial match (some sheets recognized,
        # some not) stays lenient - only a TOTAL miss is a hard error.
        raise ValueError(
            f"'{source_file}' - none of its {len(wb.sheetnames)} sheet(s) matched any known "
            f"tab format. This usually means the wrong file was passed (e.g. a position "
            f"snapshot given as --transactions). Not loading - check the file and retry."
        )

    conn.commit()
    print(f"[transactions] loaded {n_loaded} rows, skipped {n_skipped_no_account} "
          f"(unresolved account)")
    if own_conn:
        conn.close()
    return n_loaded


def _insert_txn(conn, account_id, rec):
    try:
        cur = conn.execute(
            """INSERT OR IGNORE INTO txn
                 (account_id, instrument_id, source_leg, trade_date, booking_date, value_date,
                  txn_type, txn_subtype, currency, quantity, price, gross_amount, fee_amount,
                  tax_amount, running_balance, fx_rate, realized_pl, raw_txn_type_label,
                  raw_description, source_reference, source_file)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                account_id, rec["instrument_id"], rec["source_leg"], rec["trade_date"],
                rec["booking_date"], rec["value_date"], rec["txn_type"], rec["txn_subtype"],
                rec["currency"], rec["quantity"], rec["price"], rec["gross_amount"],
                rec["fee_amount"], rec["tax_amount"], rec["running_balance"], rec["fx_rate"],
                rec["realized_pl"], rec["raw_txn_type_label"], rec["raw_description"],
                rec["source_reference"], rec["source_file"],
            ),
        )
        return cur.rowcount  # 0 if INSERT OR IGNORE silently skipped a duplicate
    except sqlite3.IntegrityError:
        return 0


if __name__ == "__main__":
    import sys
    force = "--force" in sys.argv
    load_transactions(Path(sys.argv[1]), force=force)
